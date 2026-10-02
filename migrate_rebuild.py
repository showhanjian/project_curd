#!/usr/bin/env python3
"""
重建 project_db 表：按 Excel '项目明细' 的 25 列结构重新设计，
DROP 旧表，CREATE 新表，再从 Excel 导入全部数据。
"""
import pymysql
import openpyxl
from datetime import datetime, date, timedelta

# ── 数据库配置 ──────────────────────────────────────────────
DB_CONFIG = {
    'host': '47.92.224.221',
    'port': 33067,
    'user': 'root',
    'password': 'Showhanjian@gmail.com',
    'database': 'workdb',
    'charset': 'utf8mb4',
}

# ── Excel 路径 ───────────────────────────────────────────────
EXCEL_PATH = '/root/.hermes/cache/documents/doc_a9aa6b872a17_Z1-日常工作.xlsx'
SHEET_NAME = '项目明细'

# Excel 日期序列号 → Python date
# Excel 起始日期是 1900-01-01（序列号 1），但 Python 的 openpyxl 用的就是这套
EXCEL_EPOCH = datetime(1899, 12, 30)  # 常见的兼容性处理

def excel_serial_to_date(n):
    """将 Excel 日期序列号转为 date 或 None"""
    if n is None:
        return None
    try:
        n = int(n)
        return (EXCEL_EPOCH + timedelta(days=n)).date()
    except (ValueError, TypeError, OverflowError):
        return None

def to_date(v):
    """将 datetime / int / float / str 统一转为 date 或 None"""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        return excel_serial_to_date(int(v))
    if isinstance(v, str):
        try:
            return datetime.strptime(v[:10], '%Y-%m-%d').date()
        except ValueError:
            return None
    return None

# ── 重建表 ───────────────────────────────────────────────────
def rebuild_table():
    conn = pymysql.connect(**DB_CONFIG)
    cur = conn.cursor()

    print('[1/3] DROP 旧表（如果存在）...')
    cur.execute("DROP TABLE IF EXISTS project_db")
    conn.commit()

    print('[2/3] CREATE 新表...')
    create_sql = """
    CREATE TABLE project_db (
        id              INT AUTO_INCREMENT PRIMARY KEY,
        project_no      VARCHAR(64)  COMMENT '项目编码',
        project_name    VARCHAR(256) COMMENT '项目名称',
        demand_type     VARCHAR(64)  COMMENT '需求分类',
        workload        DECIMAL(10,2) COMMENT '工作量',
        extra_work      DECIMAL(10,2) COMMENT '补记',
        remark          VARCHAR(256) COMMENT '备注',
        sign_status     VARCHAR(32)  COMMENT '签约状态',
        sign_amt        DECIMAL(14,2) COMMENT '签约金额',
        sign_date       DATE         COMMENT '签约日期',
        sign_followup   VARCHAR(256) COMMENT '签约跟踪',
        attend_start    DATE         COMMENT '考勤开始',
        attend_end      DATE         COMMENT '考勤结束',
        accept_status   VARCHAR(32)  COMMENT '验收状态',
        accept_amt      DECIMAL(14,2) COMMENT '验收金额',
        accept_date     DATE         COMMENT '验收日期',
        actual_handover DATE         COMMENT '实际上会',
        accept_followup VARCHAR(256) COMMENT '验收跟踪',
        first_pay_status VARCHAR(32)  COMMENT '首付状态',
        first_pay_amt   DECIMAL(14,2) COMMENT '首付金额',
        first_pay_date  DATE         COMMENT '首付日期',
        final_pay_status VARCHAR(32)  COMMENT '尾款状态',
        final_pay_amt   DECIMAL(14,2) COMMENT '尾款金额',
        final_pay_date  DATE         COMMENT '尾款日期',
        profit_rate     DECIMAL(10,4) COMMENT '毛利率',
        profit_amt      DECIMAL(14,2) COMMENT '毛利额'
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='项目明细'
    """
    cur.execute(create_sql)
    conn.commit()
    print('    表创建完成。')
    cur.close()
    conn.close()

# ── 读取 Excel ───────────────────────────────────────────────
def load_excel():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[SHEET_NAME]

    headers = [cell.value for cell in ws[1]]
    print(f'\n[3/3] 从 Excel 读取数据：{ws.max_row - 1} 行 × {len(headers)} 列')

    rows = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 1):
        if all(v is None for v in row):
            continue  # 跳过空行
        rows.append((i, dict(zip(headers, row))))

    print(f'    有效数据行：{len(rows)}')
    return rows

# ── 导入数据库 ───────────────────────────────────────────────
def import_data(excel_rows):
    conn = pymysql.connect(**DB_CONFIG)
    cur = conn.cursor()

    insert_sql = """
    INSERT INTO project_db (
        project_no, project_name, demand_type, workload, extra_work,
        remark, sign_status, sign_amt, sign_date, sign_followup,
        attend_start, attend_end, accept_status, accept_amt,
        accept_date, actual_handover, accept_followup,
        first_pay_status, first_pay_amt, first_pay_date,
        final_pay_status, final_pay_amt, final_pay_date,
        profit_rate, profit_amt
    ) VALUES (
        %s,%s,%s,%s,%s, %s,%s,%s,%s,%s, %s,%s,%s,%s,%s, %s,%s,%s,%s,%s, %s,%s,%s,%s,%s
    )
    """

    for excel_row_idx, rec in excel_rows:
        try:
            values = (
                str(rec.get('项目编码') or '').strip() or None,
                str(rec.get('项目名称') or '').strip() or None,
                str(rec.get('需求分类') or '').strip() or None,
                rec.get('工作量'),
                rec.get('补记'),
                str(rec.get('备注') or '').strip() or None,
                str(rec.get('签约状态') or '').strip() or None,
                rec.get('签约金额'),
                to_date(rec.get('签约日期')),
                str(rec.get('签约跟踪') or '').strip() or None,
                to_date(rec.get('考勤开始')),
                to_date(rec.get('考勤结束')),
                str(rec.get('验收状态') or '').strip() or None,
                rec.get('验收金额'),
                to_date(rec.get('验收日期')),
                to_date(rec.get('实际上会')),
                str(rec.get('验收跟踪') or '').strip() or None,
                str(rec.get('首付状态') or '').strip() or None,
                rec.get('首付金额'),
                to_date(rec.get('首付日期')),
                str(rec.get('尾款状态') or '').strip() or None,
                rec.get('尾款金额'),
                to_date(rec.get('尾款日期')),
                rec.get('毛利率'),
                rec.get('毛利额'),
            )
            cur.execute(insert_sql, values)
        except Exception as e:
            print(f'    ⚠ 行 {excel_row_idx} 导入失败：{e}')

    conn.commit()
    cur.execute('SELECT COUNT(*) FROM project_db')
    count = cur.fetchone()[0]
    print(f'    导入完成，共 {count} 条记录。')
    cur.close()
    conn.close()

# ── 主流程 ───────────────────────────────────────────────────
if __name__ == '__main__':
    rebuild_table()
    excel_rows = load_excel()
    import_data(excel_rows)
    print('\n✅ 迁移完成！')
